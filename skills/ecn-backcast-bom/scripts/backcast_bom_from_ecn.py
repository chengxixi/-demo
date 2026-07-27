import argparse
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CN_NUM = {
    1: "一",
    2: "二",
    3: "三",
    4: "四",
    5: "五",
    6: "六",
    7: "七",
    8: "八",
    9: "九",
    10: "十",
    11: "十一",
    12: "十二",
}
GREEN = "D9EAD3"
ORANGE = "FCE5CD"
RED = "F4CCCC"
BLUE = "D9EAF7"
OVERVIEW = "D9EAD3"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def cn_ordinal(index):
    return f"第{CN_NUM.get(index, index)}次"


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str) and value.strip():
        for fmt in ("%Y/%m/%d", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def date_text(value):
    return "" if value is None else f"{value.year}/{value.month}/{value.day}"


def version_key(text):
    match = re.search(r"_(V\d+\.\d+)", str(text or ""))
    return match.group(1) if match else ""


def row_is_after(row):
    return str(row.get("变更标识") or "") in {"变更后", "新增"} or str(row.get("ECN行类型") or "") == "新增子项"


def row_is_before(row):
    return str(row.get("变更标识") or "") in {"变更前", "删除"} or str(row.get("ECN行类型") or "") == "删除子项"


def diff_type(row):
    flag = str(row.get("变更标识") or "")
    row_type = str(row.get("ECN行类型") or "")
    if flag == "新增" or row_type == "新增子项":
        return "本次新增"
    if flag == "变更后":
        return "本次变更后"
    if flag == "变更前":
        return "本次变更前"
    if flag == "删除" or row_type == "删除子项":
        return "本次删除"
    return flag or row_type or "本次变更"


def row_desc(row):
    return (
        f"{row.get('子项物料编码') or ''} {row.get('子项物料名称') or ''} "
        f"用量 {row.get('用量:分子') or ''}/{row.get('用量:分母') or ''}"
    ).strip()


def normalize_bom_row(values, headers):
    row = list(values)
    if len(row) >= len(headers) + 1 and row[16] is None:
        row = row[:16] + row[17 : len(headers) + 1]
    row = row[: len(headers)] + [None] * max(0, len(headers) - len(row))
    return dict(zip(headers, row))


def read_bom_data(path, versions):
    wb = load_workbook(path, data_only=True)
    header_sheet = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [header_sheet.cell(1, c).value for c in range(1, header_sheet.max_column + 1)]
    headers = [h for h in headers if h is not None]
    lookup = {}
    version_rows = {version: [] for version in versions}
    order = {}
    order_counter = 0

    for ws in wb.worksheets:
        start = 2 if ws.cell(1, 1).value == "使用组织" else 1
        for r in range(start, ws.max_row + 1):
            values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if not any(value is not None for value in values):
                continue
            row = normalize_bom_row(values, headers)
            code = row.get("子项物料编码")
            if not code:
                continue
            code = str(code)
            lookup.setdefault(code, deepcopy(row))
            order.setdefault(code, order_counter)
            order_counter += 1
            if ws.title in version_rows:
                version_rows[ws.title].append(row)
    return headers, version_rows, lookup, order


def read_ecns_by_version(path, ecn_sheet, parent_code, versions):
    wb = load_workbook(path, data_only=True)
    ws = wb[ecn_sheet] if ecn_sheet else wb[wb.sheetnames[-1]]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    groups = []
    current = None

    for r in range(2, ws.max_row + 1):
        raw = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
        maybe_doc = raw.get("单据编号")
        if isinstance(maybe_doc, str) and maybe_doc.startswith("ECN"):
            current = {
                "doc": maybe_doc,
                "effect_date": as_date(raw.get("生效日期")),
                "change_type": raw.get("变更类型"),
                "change_method": raw.get("变更方式"),
                "reason": raw.get("原因"),
                "remark": raw.get("备注"),
                "rows": [],
            }
            groups.append(current)
        if current is None:
            continue
        bom_version = raw.get("BOM版本")
        vk = version_key(bom_version)
        if vk in versions and parent_code in str(bom_version or ""):
            raw["_version_key"] = vk
            current["rows"].append(raw)

    by_version = {version: [] for version in versions}
    for group in groups:
        for version in versions:
            rows = [row for row in group["rows"] if row["_version_key"] == version]
            if rows:
                item = deepcopy(group)
                item["rows"] = rows
                item["bom_version"] = rows[0].get("BOM版本")
                by_version[version].append(item)

    for version in versions:
        by_version[version].sort(key=lambda item: (item["effect_date"] or datetime.max.date(), item["doc"]))
    return by_version


def row_index(rows):
    return {str(row.get("子项物料编码")): i for i, row in enumerate(rows) if row.get("子项物料编码")}


def is_old_material_policy(ecn):
    return ecn.get("change_type") == "用完旧料"


def make_row_from_ecn(ecn_row, template, lookup):
    code = str(ecn_row.get("子项物料编码") or "")
    row = deepcopy(lookup.get(code) or template)
    for key in list(row.keys())[:10]:
        row[key] = None
    row["子项单位"] = "Pcs"
    row["供应组织"] = None
    row["用量:分子"] = ecn_row.get("用量:分子")
    row["用量:分母"] = ecn_row.get("用量:分母")
    row["子项物料编码"] = code
    row["子项物料名称"] = ecn_row.get("子项物料名称")
    row["子项物料禁用状态"] = row.get("子项物料禁用状态") or "否"
    return row


def ecn_applied_in_current(current_rows, ecn):
    codes = set(row_index(current_rows).keys())
    after_codes = [str(row.get("子项物料编码") or "") for row in ecn["rows"] if row_is_after(row)]
    before_codes = [str(row.get("子项物料编码") or "") for row in ecn["rows"] if row_is_before(row)]
    if after_codes:
        return all(code in codes for code in after_codes if code)
    if before_codes:
        return True if is_old_material_policy(ecn) else all(code not in codes for code in before_codes if code)
    return True


def reverse_one(rows, ecn, lookup):
    rows = deepcopy(rows)
    removed_positions = {}
    for item in ecn["rows"]:
        if row_is_after(item):
            code = str(item.get("子项物料编码") or "")
            name = str(item.get("子项物料名称") or "")
            idx = row_index(rows).get(code)
            if idx is not None:
                removed_positions.setdefault(name, idx)
                rows.pop(idx)
    for item in ecn["rows"]:
        if row_is_before(item):
            code = str(item.get("子项物料编码") or "")
            name = str(item.get("子项物料名称") or "")
            if code and code not in row_index(rows):
                rows.insert(min(removed_positions.get(name, len(rows)), len(rows)), make_row_from_ecn(item, rows[0], lookup))
    return rows


def forward_one(rows, ecn, lookup):
    rows = deepcopy(rows)
    keep_old = is_old_material_policy(ecn)
    for item in ecn["rows"]:
        if row_is_before(item) and not keep_old:
            code = str(item.get("子项物料编码") or "")
            idx = row_index(rows).get(code)
            if idx is not None:
                rows.pop(idx)
    for item in ecn["rows"]:
        if row_is_after(item):
            code = str(item.get("子项物料编码") or "")
            if code and code not in row_index(rows):
                rows.append(make_row_from_ecn(item, rows[0], lookup))
    return rows


def build_snapshot(current_rows, ecns, applied, ordinal, lookup):
    state = deepcopy(current_rows)
    for j in range(len(ecns), ordinal, -1):
        if applied[j - 1]:
            state = reverse_one(state, ecns[j - 1], lookup)
    for j in range(1, ordinal + 1):
        if not applied[j - 1]:
            state = forward_one(state, ecns[j - 1], lookup)
    return state


def label_rows(rows, headers, label, ecn, valid_until):
    rows = deepcopy(rows)
    parent = deepcopy(rows[0]) if rows else {}
    for i, row in enumerate(rows, start=1):
        row["项次"] = i
        if i == 1:
            for key in headers[:10]:
                row[key] = row.get(key) or parent.get(key)
        else:
            for key in headers[:10]:
                row[key] = None
        row["_倒推版本标签"] = label
        row["_对应ECN"] = ecn["doc"]
        row["_BOM版本"] = ecn.get("bom_version") or ""
        row["_ECN生效日期"] = date_text(ecn.get("effect_date"))
        row["_最后有效时间"] = date_text(valid_until) if valid_until else "当前有效"
        row["_本次变更标记"] = ""
        row["_本次变更备注"] = ""
    return rows


def style_header(row, fill):
    for cell in row:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_row(ws, r, fill, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(r, c)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def pair_changes(ecn_rows):
    before_by_name = {}
    for row in ecn_rows:
        if row_is_before(row):
            before_by_name.setdefault(str(row.get("子项物料名称") or "").strip(), []).append(row)
    pair_by_after_code = {}
    paired_before_ids = set()
    for row in ecn_rows:
        if not row_is_after(row):
            continue
        name = str(row.get("子项物料名称") or "").strip()
        before = before_by_name.get(name, [])
        if before:
            item = before.pop(0)
            pair_by_after_code[str(row.get("子项物料编码") or "")] = item
            paired_before_ids.add(id(item))
    pure_deletes = [row for row in ecn_rows if row_is_before(row) and id(row) not in paired_before_ids]
    return pair_by_after_code, pure_deletes


def add_delete_note_far(ws, pure_deletes):
    if not pure_deletes:
        return 0
    start = ws.max_row + 6
    labels = ["本次仅删除物料说明（不属于本页有效BOM，不参与BOM行数）", "原料号", "物料名称", "用量", "备注"]
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(start, c)
        cell.value = label
        cell.font = Font(bold=True, color="9C0006" if c == 1 else None)
        cell.fill = PatternFill("solid", fgColor=RED)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(pure_deletes, start=1):
        r = start + i
        ws.cell(r, 2).value = row.get("子项物料编码")
        ws.cell(r, 3).value = row.get("子项物料名称")
        ws.cell(r, 4).value = f"{row.get('用量:分子') or ''}/{row.get('用量:分母') or ''}"
        ws.cell(r, 5).value = f"{diff_type(row)}：本次ECN只删除该物料，无对应新增料号。{row_desc(row)}"
        for c in range(1, 6):
            cell = ws.cell(r, c)
            cell.fill = PatternFill("solid", fgColor=RED)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions["E"].width = max(ws.column_dimensions["E"].width or 10, 74)
    return len(pure_deletes)


def write_bom_sheet(ws, headers, rows, ecn):
    extra = ["倒推版本标签", "对应ECN", "BOM版本", "ECN生效日期", "最后有效时间", "本次变更标记", "本次变更备注"]
    ws.append(headers + extra)
    for row in rows:
        ws.append([row.get(h) for h in headers] + [row.get("_" + h) for h in extra])
    style_header(ws[1], BLUE)
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14 if col < 16 else 22
    ws.column_dimensions["S"].width = 52
    for col in range(len(headers) + 1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 28
    ws.column_dimensions[get_column_letter(ws.max_column)].width = 82
    mark_current_ecn(ws, ecn)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def mark_current_ecn(ws, ecn):
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    if "子项物料编码" not in headers:
        return
    code_col = headers.index("子项物料编码") + 1
    mark_col = headers.index("本次变更标记") + 1
    remark_col = headers.index("本次变更备注") + 1
    idx = {str(ws.cell(r, code_col).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, code_col).value}
    pair_by_after_code, pure_deletes = pair_changes(ecn["rows"])
    max_col = ws.max_column

    for row in ecn["rows"]:
        if not row_is_after(row):
            continue
        code = str(row.get("子项物料编码") or "")
        r = idx.get(code)
        if not r:
            continue
        mark = diff_type(row)
        old = pair_by_after_code.get(code)
        remark = f"{mark}：本页对应ECN {ecn['doc']} 的变动物料。{row_desc(row)}"
        if old:
            remark += (
                f"；由原料号 {old.get('子项物料编码')}（{old.get('子项物料名称')}，"
                f"用量 {old.get('用量:分子')}/{old.get('用量:分母')}）变更/替换而来。"
            )
        ws.cell(r, mark_col).value = mark
        ws.cell(r, remark_col).value = remark
        style_row(ws, r, GREEN, max_col)

    if ecn.get("change_type") == "用完旧料":
        for old in pair_by_after_code.values():
            code = str(old.get("子项物料编码") or "")
            r = idx.get(code)
            if r:
                ws.cell(r, mark_col).value = "用完旧料保留"
                ws.cell(r, remark_col).value = f"用完旧料保留：本页对应ECN {ecn['doc']} 的原料号仍保留在BOM中。{row_desc(old)}"
                style_row(ws, r, ORANGE, max_col)

    add_delete_note_far(ws, pure_deletes)


def changes_text(ecn):
    return "；".join(
        f"{row.get('变更标识') or row.get('ECN行类型')} {row.get('子项物料编码')} {row.get('子项物料名称')} {row.get('用量:分子') or ''}/{row.get('用量:分母') or ''}".strip()
        for row in ecn["rows"]
    )


def safe_sheet_title(version, index, valid_until):
    suffix = f"至{valid_until.month}.{valid_until.day}" if valid_until else "当前有效"
    return f"{version}第{index}次_{suffix}"[:31]


def write_overview(wb, parent_code, version, ecns):
    ws = wb.active
    ws.title = f"{version}总览"
    ws.append(["产品编码", parent_code, "版本", version])
    ws.append([
        "说明",
        "每个明细页是一整张BOM；每页只标记该页对应ECN的变动物料；有新增料号时，原料号写入新增料号备注；仅删除无新增时，放在远离BOM的说明区。"
        if ecns
        else "该版本在ECN附件中没有记录。",
    ])
    ws.append([])
    headers = ["序号", "版本标签", "对应ECN", "BOM版本", "ECN生效日期", "最后有效时间", "ECN行数", "BOM行数", "当前BOM是否已反映", "变更内容"]
    ws.append(headers)
    style_header(ws[4], OVERVIEW)
    widths = [8, 42, 18, 30, 14, 16, 10, 10, 18, 90]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"
    if ecns:
        ws.auto_filter.ref = f"A4:J{4 + len(ecns)}"
    return ws


def generate_version(out_dir, parent_code, version, headers, current_rows, lookup, ecns):
    out_path = out_dir / f"{parent_code}_{version}_BOM倒推明细_明细内标记差异.xlsx"
    wb = Workbook()
    overview = write_overview(wb, parent_code, version, ecns)

    if not ecns:
        rows = deepcopy(current_rows)
        for i, row in enumerate(rows, start=1):
            row["项次"] = i
        ws = wb.create_sheet(f"{version}当前BOM")
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
        style_header(ws[1], BLUE)
        wb.save(out_path)
        return out_path

    applied = [ecn_applied_in_current(current_rows, ecn) for ecn in ecns]
    for i, ecn in enumerate(ecns, start=1):
        next_effect = ecns[i]["effect_date"] if i < len(ecns) else None
        valid_until = next_effect - timedelta(days=1) if next_effect else None
        label = f"{version}({cn_ordinal(i)}变更，最后有效时间：{date_text(valid_until) if valid_until else '当前有效'})"
        snapshot = build_snapshot(current_rows, ecns, applied, i, lookup)
        rows = label_rows(snapshot, headers, label, ecn, valid_until)
        ws = wb.create_sheet(safe_sheet_title(version, i, valid_until))
        write_bom_sheet(ws, headers, rows, ecn)
        overview.append([
            i,
            label,
            ecn["doc"],
            ecn.get("bom_version") or "",
            date_text(ecn.get("effect_date")),
            date_text(valid_until) if valid_until else "当前有效",
            len(ecn["rows"]),
            len(snapshot),
            "是" if applied[i - 1] else "否",
            changes_text(ecn),
        ])

    wb.save(out_path)
    return out_path


def main():
    parser = argparse.ArgumentParser(description="Backcast BOM snapshots from ECN records.")
    parser.add_argument("--bom", required=True, help="BOM workbook path")
    parser.add_argument("--ecn", required=True, help="ECN workbook path")
    parser.add_argument("--parent-code", required=True, help="Target parent material code")
    parser.add_argument("--ecn-sheet", default="", help="ECN sheet name; defaults to the last sheet")
    parser.add_argument("--versions", nargs="+", required=True, help="BOM versions to output, e.g. V1.0 V1.1")
    parser.add_argument("--out-dir", required=True, help="Output directory")
    parser.add_argument("--zip", action="store_true", help="Create a zip archive of generated workbooks")
    args = parser.parse_args()

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    headers, version_rows, lookup, _ = read_bom_data(args.bom, args.versions)
    ecns_by_version = read_ecns_by_version(args.ecn, args.ecn_sheet, args.parent_code, args.versions)

    outputs = []
    for version in args.versions:
        outputs.append(generate_version(out_dir, args.parent_code, version, headers, version_rows[version], lookup, ecns_by_version[version]))

    if args.zip:
        zip_path = out_dir / f"{args.parent_code}_BOM倒推明细_明细内标记差异.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in outputs:
                archive.write(path, path.name)
        outputs.append(zip_path)

    print("生成完成:")
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
